"""fz-quiz: a Quiz for a Topic, as an Excel workbook in the LMS Quiz Template.

The user's choices become a Quiz Request; the builder plans the question slots (type, difficulty,
points, Video) so the Host App only writes each question, and every LMS import rule is checked in code.
Cells keep their Fuzentrix markup (`code`, ``` blocks, **bold**), so fz-quiz-html can turn them into the
LMS's RICH_HTML without guessing; every field is checked against its LMS HTML context when submitted.
"""
from __future__ import annotations

from .bundle import read_json, write_json
from .problems import Rejected

from .vocab import CODE_QUESTIONS, DIFFICULTIES, EXPLANATIONS, QUESTION_TYPES as TYPES  # noqa: F401
MAX_QUESTIONS = 100  # LMS import limit per file
DEFAULT_DIFFICULTY_SHARE = {"easy": 0.4, "medium": 0.4, "hard": 0.2}
DEFAULT_TYPE_SHARE = {"SINGLE_SELECT": 0.6, "MULTI_SELECT": 0.2, "TRUE_FALSE": 0.2}
DEFAULT_POINTS = {"easy": 1, "medium": 1, "hard": 2}


def split(total: int, shares: dict) -> dict:
    """Largest-remainder split of `total` by `shares` (keys keep their order)."""
    raw = {k: total * v for k, v in shares.items()}
    out = {k: int(v) for k, v in raw.items()}
    for k in sorted(raw, key=lambda k: raw[k] - out[k], reverse=True)[: total - sum(out.values())]:
        out[k] += 1
    return out


# --- Quiz Request -------------------------------------------------------------

def submit_request(bundle, data: dict) -> dict:
    problems = []
    total = data.get("questions_total", 20)
    if not isinstance(total, int) or total < 1:
        raise Rejected("questions_total must be a whole number of 1 or more.")
    if total > MAX_QUESTIONS:
        raise Rejected(f"questions_total is {total}; the LMS imports at most {MAX_QUESTIONS} questions per file. Split the Topic into two Quizzes.")
    by_difficulty = data.get("by_difficulty") or split(total, DEFAULT_DIFFICULTY_SHARE)
    by_type = data.get("by_type") or split(total, DEFAULT_TYPE_SHARE)
    for name, counts, allowed in (("by_difficulty", by_difficulty, DIFFICULTIES), ("by_type", by_type, TYPES)):
        bad = [k for k in counts if k not in allowed]
        if bad:
            problems.append(f"{name}: {', '.join(bad)} not one of {', '.join(allowed)}.")
        if any(not isinstance(v, int) or v < 0 for v in counts.values()):
            problems.append(f"{name}: counts must be whole numbers of 0 or more.")
        elif sum(counts.values()) != total:
            problems.append(f"{name} adds up to {sum(counts.values())}, but questions_total is {total}.")
    points = dict(DEFAULT_POINTS, **(data.get("points") or {}))
    if any(k not in DIFFICULTIES or not isinstance(v, int) or v < 1 for k, v in points.items()):
        problems.append("points: give a whole number of 1 or more for easy, medium and hard.")
    options = data.get("options_per_question", 4)
    if options not in (2, 3, 4):
        problems.append("options_per_question must be 2, 3 or 4 (the LMS template has four option columns).")
    explanations = data.get("explanations", "correct_only")
    if explanations not in EXPLANATIONS:
        problems.append(f"explanations must be one of {', '.join(EXPLANATIONS)}.")
    code_questions = data.get("code_questions", "auto")
    if code_questions not in CODE_QUESTIONS:
        problems.append(f"code_questions must be one of {', '.join(CODE_QUESTIONS)}.")
    if problems:
        raise Rejected(problems)

    request = {
        "questions_total": total,
        "by_difficulty": {k: by_difficulty.get(k, 0) for k in DIFFICULTIES},
        "by_type": {k: by_type.get(k, 0) for k in TYPES},
        "points": points,
        "options_per_question": options,
        "explanations": explanations,
        "code_questions": code_questions,
        "notes": data.get("notes", ""),
    }
    plan = plan_slots(bundle, request)
    write_json(bundle.path("quiz", "request.json"), request)
    write_json(bundle.path("quiz", "plan.json"), plan)
    bundle.manifest["stages"]["quiz_request"] = {"status": "done"}
    bundle.manifest["stages"].pop("quiz", None)
    return request


def plan_slots(bundle, request: dict) -> list:
    """Pair difficulties with types (easy ↔ TRUE_FALSE first), then deal the slots to Videos in turn,
    each Video's share proportional to how much is said in it."""
    from .topic import masters

    difficulties = [d for d in DIFFICULTIES for _ in range(request["by_difficulty"][d])]
    types = [t for t in ("TRUE_FALSE", "SINGLE_SELECT", "MULTI_SELECT") for _ in range(request["by_type"][t])]
    words = {vid: _spoken_words(m) for vid, m in masters(bundle).items()}
    total_words = sum(words.values()) or 1
    quota = split(request["questions_total"], {vid: w / total_words for vid, w in words.items()})
    videos = list(words)
    slots, turn = [], 0
    for i, (difficulty, qtype) in enumerate(zip(difficulties, types), 1):
        while quota[videos[turn % len(videos)]] == 0:
            turn += 1
        vid = videos[turn % len(videos)]
        quota[vid] -= 1
        turn += 1
        slots.append({"id": f"Q{i:02d}", "video": vid, "type": qtype, "difficulty": difficulty,
                      "points": request["points"][difficulty]})
    return slots


def _spoken_words(master: dict) -> int:
    return sum(len(s["text"].split()) for seg in master["segments"] for s in seg.get("spoken", []))


def draft_video(bundle, video_id: str) -> dict:
    if "quiz_request" not in bundle.manifest["stages"]:
        raise Rejected("Submit the Quiz Request first (`fz.py submit <bundle> quiz-request <file>`).")
    if video_id not in bundle.manifest["videos"]:
        raise Rejected(f"{video_id} is not in this Topic ({', '.join(bundle.manifest['videos'])}).")
    request = read_json(bundle.path("quiz", "request.json"))
    questions = []
    for slot in read_json(bundle.path("quiz", "plan.json")):
        if slot["video"] != video_id:
            continue
        if slot["type"] == "TRUE_FALSE":
            options = [{"text": "True", "correct": None, "explanation": ""}, {"text": "False", "correct": None, "explanation": ""}]
        else:
            options = [{"text": "", "correct": None, "explanation": ""} for _ in range(request["options_per_question"])]
        questions.append({"id": slot["id"], "type": slot["type"], "difficulty": slot["difficulty"], "points": slot["points"],
                          "source_segments": [], "question": "", "options": options, "answer_check": ""})
    return {"video": video_id, "questions": questions}


# --- submitted questions ---------------------------------------------------------------

import re  # noqa: E402

LETTERS = "ABCD"
_ABOVE = re.compile(r"^\s*(all|none)\s+of\s+the\s+above\.?\s*$", re.IGNORECASE)


def written_videos(bundle) -> list:
    folder = bundle.path("quiz", "questions")
    return sorted(p.stem for p in folder.glob("*.json")) if folder.exists() else []


def submit_questions(bundle, data: dict) -> str:
    from .topic import masters

    if "quiz_request" not in bundle.manifest["stages"]:
        raise Rejected("Submit the Quiz Request first (`fz.py submit <bundle> quiz-request <file>`).")
    vid = data.get("video")
    if vid not in bundle.manifest["videos"]:
        raise Rejected(f"'video' must be one of this Topic's Videos ({', '.join(bundle.manifest['videos'])}).")
    request = read_json(bundle.path("quiz", "request.json"))
    slots = {s["id"]: s for s in read_json(bundle.path("quiz", "plan.json")) if s["video"] == vid}
    segments = {s["id"] for s in masters(bundle)[vid]["segments"]}
    others = {q["question"].strip().lower(): q["id"]
              for other in written_videos(bundle) if other != vid
              for q in read_json(bundle.path("quiz", "questions", f"{other}.json"))["questions"]}
    questions = data.get("questions") or []
    given = [q.get("id") for q in questions]
    problems = [f"{sid} is missing (planned for {vid})." for sid in slots if sid not in given]
    problems += [f"{qid} is not a planned question for {vid}." for qid in given if qid not in slots]
    seen = {}
    for q in questions:
        if q.get("id") in slots:
            problems += _check_question(q, slots[q["id"]], request, segments, vid)
            key = str(q.get("question", "")).strip().lower()
            if key and (key in seen or key in others):
                problems.append(f"{q['id']} repeats {seen.get(key) or others[key]}.")
            seen[key] = q["id"]
    problems += _html_problems(bundle, [q for q in questions if q.get("id") in slots])
    if problems:
        raise Rejected(problems)
    write_json(bundle.path("quiz", "questions", f"{vid}.json"), data)
    bundle.manifest["stages"]["quiz"] = {"status": "in_progress"}
    bundle.manifest["gates"].get("3", {}).pop("quiz", None)
    return vid


def _html_problems(bundle, questions: list) -> list:
    """Each field, as fz-quiz-html will write it, against its LMS HTML context: found now, not at upload."""
    from . import lmspolicy, markup

    policy, problems = lmspolicy.load(bundle), []
    for q in questions:
        fields = [("question", q.get("question"), "quiz_question")]
        for i, o in enumerate(q.get("options") or []):
            fields += [(f"option {LETTERS[i] if i < 4 else i + 1}", o.get("text"), "quiz_option"),
                       (f"option {LETTERS[i] if i < 4 else i + 1} explanation", o.get("explanation"), "quiz_explanation")]
        for label, text, context in fields:
            if isinstance(text, str) and text.strip():
                problems += [f"{q['id']} {label}: {p}" for p in lmspolicy.check(markup.render(text, policy, context), policy, context)]
    return problems


def _check_question(q: dict, slot: dict, request: dict, segments: set, vid: str) -> list:
    qid, qtype = q["id"], slot["type"]
    problems = []
    for key in ("type", "difficulty", "points"):
        if q.get(key) != slot[key]:
            problems.append(f"{qid}: {key} must be {slot[key]} (as planned).")
    if not str(q.get("question", "")).strip():
        problems.append(f"{qid}: 'question' is empty.")
    if not str(q.get("answer_check", "")).strip():
        problems.append(f"{qid}: 'answer_check' is empty: re-derive the correct answer from the Video before marking it.")
    sources = q.get("source_segments") or []
    if not sources:
        problems.append(f"{qid}: 'source_segments' must list the Segment(s) the question comes from.")
    for sid in sources:
        if sid not in segments:
            problems.append(f"{qid}: source segment {sid} is not in {vid}.")
    options = q.get("options") or []
    texts = [str(o.get("text", "")).strip() for o in options]
    if qtype == "TRUE_FALSE":
        if texts != ["True", "False"]:
            problems.append(f"{qid}: TRUE_FALSE options must be exactly 'True' and 'False', in that order.")
    elif len(options) != request["options_per_question"]:
        problems.append(f"{qid}: needs {request['options_per_question']} options (has {len(options)}).")
    for i, o in enumerate(options):
        letter = LETTERS[i] if i < len(LETTERS) else str(i + 1)
        if not texts[i]:
            problems.append(f"{qid}: option {letter} is empty.")
        if _ABOVE.match(texts[i]):
            problems.append(f"{qid}: option {letter}: 'All of the above' / 'None of the above' breaks when options are reordered; write a real option.")
        if not isinstance(o.get("correct"), bool):
            problems.append(f"{qid}: option {letter}: set 'correct' to true or false.")
        explained = str(o.get("explanation", "")).strip()
        if o.get("correct") is True and not explained:
            problems.append(f"{qid}: option {letter} is correct but has no explanation.")
        elif request["explanations"] == "all_options" and not explained:
            problems.append(f"{qid}: option {letter} has no explanation.")
    lowered = [t.lower() for t in texts if t]
    if len(set(lowered)) != len(lowered):
        problems.append(f"{qid}: two options have the same text.")
    n_correct = sum(1 for o in options if o.get("correct") is True)
    if qtype in ("SINGLE_SELECT", "TRUE_FALSE") and n_correct != 1:
        problems.append(f"{qid}: {qtype} needs exactly 1 correct option (has {n_correct}).")
    if qtype == "MULTI_SELECT" and not 1 <= n_correct < len(options):
        problems.append(f"{qid}: MULTI_SELECT needs at least 1 correct option and at least 1 wrong one (has {n_correct} of {len(options)} correct).")
    return problems


# --- building the LMS workbook ----------------------------------------------------------

from copy import copy as _copy  # noqa: E402
from pathlib import Path as _Path  # noqa: E402

TEMPLATE = _Path(__file__).parent / "assets" / "lms_quiz_template.xlsx"


def to_plain(text: str) -> str:
    from .markup import unmark

    return unmark(text)


def has_markup(questions: list) -> bool:
    texts = [q["question"] for q in questions] + [t for q in questions for o in q["options"] for t in (o["text"], o.get("explanation") or "")]
    return any("`" in t or "**" in t for t in texts)


def _arranged(bundle) -> list:
    """All questions in plan order, with options reordered so correct answers are spread across A–D."""
    plan = read_json(bundle.path("quiz", "plan.json"))
    by_id = {q["id"]: q for vid in written_videos(bundle) for q in read_json(bundle.path("quiz", "questions", f"{vid}.json"))["questions"]}
    counters, out = {"SINGLE_SELECT": 0, "MULTI_SELECT": 0}, []
    for slot in plan:
        q = dict(by_id[slot["id"]], video=slot["video"])
        options = list(q["options"])
        if q["type"] == "SINGLE_SELECT":
            correct = next(o for o in options if o["correct"])
            rest = [o for o in options if o is not correct]
            target = counters["SINGLE_SELECT"] % len(options)
            options = rest[:target] + [correct] + rest[target:]
            counters["SINGLE_SELECT"] += 1
        elif q["type"] == "MULTI_SELECT":
            k = counters["MULTI_SELECT"] % len(options)
            options = options[-k:] + options[:-k] if k else options
            counters["MULTI_SELECT"] += 1
        q["options"] = options
        out.append(q)
    return out


def lms_problems(questions: list) -> list:
    """The LMS Rules sheet, checked on the final rows (all-or-nothing import)."""
    problems = []
    if len(questions) > MAX_QUESTIONS:
        problems.append(f"{len(questions)} questions; the LMS imports at most {MAX_QUESTIONS}.")
    for q in questions:
        n = len(q["options"])
        correct = sum(1 for o in q["options"] if o["correct"])
        if not q["question"].strip() or q["type"] not in TYPES or not isinstance(q["points"], int) or q["points"] < 1:
            problems.append(f"{q['id']}: question, type or points not valid for the LMS.")
        if n < 2 or n > 4:
            problems.append(f"{q['id']}: needs 2–4 options for the LMS.")
        if q["type"] in ("SINGLE_SELECT", "TRUE_FALSE") and correct != 1:
            problems.append(f"{q['id']}: exactly 1 option must be TRUE.")
        if q["type"] == "MULTI_SELECT" and correct < 1:
            problems.append(f"{q['id']}: at least 1 option must be TRUE.")
        if q["type"] == "TRUE_FALSE" and n != 2:
            problems.append(f"{q['id']}: TRUE_FALSE uses options A and B only.")
    return problems


def template(bundle):
    """The LMS Quiz Template: the Configuration Pack's copy when it has one, else the skill's."""
    pack = bundle.path("_config_snapshot", "templates", "lms_quiz_template.xlsx")
    return pack if pack.exists() else TEMPLATE


def _write_workbook(path, questions: list, template_path=TEMPLATE) -> None:
    """The LMS Quiz Template filled in, cells in Fuzentrix markup, marked as ours for fz-quiz-html."""
    from openpyxl import load_workbook

    from .quizhtml import MARKER

    wb = load_workbook(template_path)
    ws = wb["Quiz Questions"]
    col = {c.value: c.column for c in ws[1]}
    proto = {c.column: c for c in ws[2]}
    height = ws.row_dimensions[2].height
    for r, q in enumerate(questions, 2):
        values = {"Question": q["question"], "Question Type": q["type"], "Points": q["points"], "Content Format": "PLAIN_TEXT"}
        for i, o in enumerate(q["options"]):
            letter = LETTERS[i]
            values[f"Option {letter}"] = o["text"]
            values[f"Option {letter} Correct"] = "TRUE" if o["correct"] else "FALSE"
            if str(o.get("explanation", "")).strip():
                values[f"Option {letter} Explanation"] = o["explanation"]
        for header, c in col.items():
            cell = ws.cell(r, c)
            cell._style = _copy(proto[c]._style)
            cell.value = values.get(header)
        ws.row_dimensions[r].height = height
    wb.properties.keywords = MARKER
    _Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build(bundle) -> dict:
    from . import views

    missing = [vid for vid in bundle.manifest["videos"] if vid not in written_videos(bundle)]
    if "quiz_request" not in bundle.manifest["stages"]:
        raise Rejected("Submit the Quiz Request first.")
    if missing:
        raise Rejected([f"{vid} has no Quiz questions yet." for vid in missing])
    questions = _arranged(bundle)
    problems = lms_problems(questions)
    if problems:
        raise Rejected(problems)
    tid = bundle.manifest["topic"]["id"]
    _write_workbook(bundle.path("quiz", f"{tid}_quiz.xlsx"), questions, template(bundle))
    views.write_pair(bundle.path("quiz", "quiz_preview"), _preview(bundle, questions), title=f"Quiz: {bundle.manifest['topic']['title'] or tid}")
    bundle.manifest["stages"]["quiz"] = {"status": "built", "questions": len(questions), "markup": has_markup(questions)}
    bundle.manifest["gates"].get("3", {}).pop("quiz", None)
    bundle.log("quiz built")
    return {"questions": len(questions), "points": sum(q["points"] for q in questions)}


def _preview(bundle, questions: list) -> str:
    from .views import config_note

    count = lambda key: ", ".join(f"{k} {sum(1 for q in questions if q[key] == k)}" for k in dict.fromkeys(q[key] for q in questions))
    positions = {l: 0 for l in LETTERS}
    for q in questions:
        if q["type"] == "SINGLE_SELECT":
            positions[LETTERS[next(i for i, o in enumerate(q["options"]) if o["correct"])]] += 1
    out = [f"# Quiz: {bundle.manifest['topic']['title'] or bundle.manifest['topic']['id']}", "", config_note(bundle), "",
           f"- Questions: {len(questions)} · Points: {sum(q['points'] for q in questions)}",
           f"- Types: {count('type')}", f"- Difficulty: {count('difficulty')}", f"- Videos: {count('video')}",
           "- Correct answer positions (single select): " + ", ".join(f"{l} {n}" for l, n in positions.items()), ""]
    for q in questions:
        out += [f"## {q['id']} · {q['type'].replace('_', ' ').lower()} · {q['difficulty']} · {q['points']} pt · {q['video']} {', '.join(q['source_segments'])}", "",
                to_plain(q["question"]), ""]
        for i, o in enumerate(q["options"]):
            mark = "**✔ correct**" if o["correct"] else ""
            out.append(f"- {LETTERS[i]}. {to_plain(o['text'])} {mark}".rstrip())
            if str(o.get("explanation", "")).strip():
                out.append(f"- _Explanation ({LETTERS[i]}):_ {to_plain(o['explanation'])}")
        out.append("")
    return "\n".join(out) + "\n"
