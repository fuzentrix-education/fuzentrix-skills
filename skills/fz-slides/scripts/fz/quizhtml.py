"""fz-quiz-html: turn a quiz workbook's questions, options and explanations into LMS RICH_HTML.

Works on the workbooks fz-quiz writes (Fuzentrix markup in the cells, so nothing is guessed) and on any
existing quiz workbook in the LMS template (code is recognised by its look). Each field is written under its
own LMS HTML Policy context (question, option, explanation) and checked; one bad cell stops the whole file,
because the LMS rejects a whole import for one bad row. An AI may mark code in an existing workbook first
(quiz-cells, then --marks): only marks are accepted; any change to the words is refused. Nothing but the
converted cells and the Content Format column changes."""
from __future__ import annotations

import re
from pathlib import Path

from . import lmspolicy, markup
from .bundle import read_json, write_json
from .problems import Rejected

MARKER = "fz-markup-1"   # set on workbooks fz-quiz writes: their cells are Fuzentrix markup
HTML_TAG = re.compile(r"</?(p|br|strong|em|b|i|u|s|code|pre|ul|ol|li|table|thead|tbody|tr|th|td|blockquote|sub|sup|h[2-6])\b[^>]*>", re.I)
CODE_LINE = re.compile(r"^\s*(def |class |for |while |if |elif |else:|return |import |from \w+ import|print\(|SELECT |INSERT |UPDATE |DELETE |"
                       r"CREATE |ALTER |DROP |#include|public |private |static |System\.out|printf\(|scanf\(|int |float |char |double |"
                       r"let |const |var |function |console\.)|[;{}]\s*$|^\s{2,}\S")
TOKEN = re.compile(r"(?<![`\w.])([A-Za-z_]\w*(?:\.[A-Za-z_]\w*)*\(\)|[a-z]+(?:_[a-z0-9]+)+)(?![\w`])")


def _fields(ws, header_row: int) -> dict:
    """column → field: quiz_question, quiz_option, quiz_explanation, or format."""
    out = {}
    for c in ws[header_row]:
        h = " ".join(str(c.value or "").split()).lower()
        if h == "content format":
            out[c.column] = "format"
        elif h in ("question", "questions", "question text", "problem", "problem statement"):
            out[c.column] = "quiz_question"
        elif re.fullmatch(r"option [a-z0-9]+ explanation", h):
            out[c.column] = "quiz_explanation"
        elif re.fullmatch(r"option [a-z0-9]+", h):
            out[c.column] = "quiz_option"
    return out


def _sheets(wb, name: str):
    """(sheet, header row, fields) for every sheet holding quiz questions."""
    found = []
    for ws in wb.worksheets:
        def is_header(r):
            names = {" ".join(str(c.value).split()).lower() for c in ws[r] if c.value}
            return bool(names & {"question", "questions", "question text"}) and any(re.fullmatch(r"option [a-z0-9]+", n) for n in names)
        header_row = next((r for r in range(1, min(ws.max_row, 25) + 1) if is_header(r)), None)
        if header_row is None:
            continue
        fields = _fields(ws, header_row)
        if "format" not in fields.values():
            raise Rejected(f"{name}: sheet '{ws.title}' has a Question column but no Content Format column, so the LMS can't be told the "
                           "cells are HTML. Use the LMS Quiz Template.")
        if list(fields.values()).count("format") > 1:
            raise Rejected(f"{name}: sheet '{ws.title}' has more than one Content Format column.")
        found.append((ws, header_row, fields))
    if not found:
        raise Rejected(f"{name} has no sheet of quiz questions (a 'Question' column with 'Option A', 'Option B'…).")
    return found


def _rows(ws, header_row: int, fields: dict):
    qcol = next(c for c, f in fields.items() if f == "quiz_question")
    for r in range(header_row + 1, ws.max_row + 1):
        if str(ws.cell(r, qcol).value or "").strip():
            yield r


def _language(lines: list) -> str:
    text = "\n".join(lines)
    for lang, pattern in (("sql", r"\b(SELECT|INSERT|UPDATE|DELETE|CREATE|ALTER|DROP)\b"), ("python", r"\bdef |print\(|elif |import "),
                          ("java", r"System\.out|public |private "), ("c", r"#include|printf\(|scanf\("), ("javascript", r"console\.|function |let |const ")):
        if re.search(pattern, text):
            return lang
    return ""


def auto_mark(text: str) -> str:
    """Fuzentrix markup for an unmarked cell from an existing workbook: a paragraph that reads as code becomes a
    code block; function calls (name()) and snake_case names become `code`, once each."""
    paras, out = re.split(r"(\n\s*\n)", text), []
    for para in paras:
        lines = para.split("\n")
        if para.strip() and len(lines) >= 2 and sum(bool(CODE_LINE.search(l)) for l in lines) >= max(2, 0.6 * len(lines)):
            out.append(f"```{_language(lines)}\n{para.strip(chr(10))}\n```")
            continue
        seen = set()

        def mark(m):
            if m.group(1) in seen:
                return m.group(1)
            seen.add(m.group(1))
            return f"`{m.group(1)}`"
        out.append(TOKEN.sub(mark, para))
    return "".join(out)


def _same_words(a: str, b: str) -> bool:
    norm = lambda t: " ".join(markup.unmark(t).split())
    return norm(a) == norm(b)


# --- the optional AI marking pass ------------------------------------------------------------

def cells(source, out_file) -> int:
    """List the workbook's content cells (sheet!cell, field, text) for an AI to add marks to."""
    from openpyxl import load_workbook

    wb = load_workbook(source)
    listed = []
    for ws, header_row, fields in _sheets(wb, Path(source).name):
        for r in _rows(ws, header_row, fields):
            for c, field in fields.items():
                value = ws.cell(r, c).value
                if field != "format" and isinstance(value, str) and value.strip() and not value.startswith("=") and not HTML_TAG.search(value):
                    listed.append({"ref": f"{ws.title}!{ws.cell(r, c).coordinate}", "field": field, "text": value})
    write_json(out_file, {"workbook": Path(source).name, "cells": listed, "marks": []})
    return len(listed)


# --- convert ------------------------------------------------------------------------------------

def convert(source, target, policy: dict, marks: dict = None, plain=None) -> dict:
    """Write target as RICH_HTML (and, if plain is given, a PLAIN_TEXT copy without marks)."""
    from openpyxl import load_workbook

    source, target = Path(source), Path(target)
    if source.resolve() == target.resolve():
        raise Rejected("Write the converted workbook to a new file; the original stays unchanged.")
    try:
        wb = load_workbook(source)
    except Exception as e:
        raise Rejected(f"{source.name} is not a readable .xlsx workbook ({e.__class__.__name__}).")
    ours = (wb.properties.keywords or "") == MARKER
    marks = {m["ref"]: m["marked"] for m in (marks or {}).get("marks") or [] if m.get("ref") and isinstance(m.get("marked"), str)}
    problems, counts = [], {"rows": 0, "cells": 0, "marked": 0, "guessed": 0, "html": 0, "sheets": 0}
    plain_values = {}
    for ws, header_row, fields in _sheets(wb, source.name):
        counts["sheets"] += 1
        fcol = next(c for c, f in fields.items() if f == "format")
        for r in _rows(ws, header_row, fields):
            counts["rows"] += 1
            for c, field in fields.items():
                cell = ws.cell(r, c)
                if field == "format" or not isinstance(cell.value, str) or not cell.value.strip() or cell.value.startswith("="):
                    continue
                ref, text = f"{ws.title}!{cell.coordinate}", cell.value
                if ref in marks:
                    marked = marks.pop(ref)
                    if not _same_words(marked, text):
                        problems.append(f"{ref}: the marked text changes the words; only add marks (`code`, **bold**, ``` blocks).")
                        continue
                    source_text = marked
                    counts["marked"] += 1
                elif ours or "`" in text:
                    source_text = text  # already Fuzentrix markup
                elif HTML_TAG.search(text):
                    source_text = None  # written as HTML already
                    counts["html"] += 1
                else:
                    source_text = auto_mark(text)
                    counts["guessed"] += source_text != text
                html_ = text.strip() if source_text is None else markup.render(source_text, policy)
                found = lmspolicy.check(html_, policy, field)
                if found:
                    problems += [f"{ref} ({field.replace('quiz_', '')}): {p}" for p in found]
                    continue
                plain_values[(ws.title, cell.coordinate)] = markup.unmark(text) if source_text is not None else text
                cell.value = html_
                counts["cells"] += 1
            ws.cell(r, fcol).value = "RICH_HTML"
    problems += [f"{ref}: no such content cell in {source.name}." for ref in marks]
    if problems:
        raise Rejected(problems + ["Nothing was written: the LMS rejects the whole file for one bad cell."])
    wb.properties.keywords = None
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    reopened = load_workbook(target)  # the LMS reads it back; so do we
    if sum(1 for ws, h, f in _sheets(reopened, target.name) for _ in _rows(ws, h, f)) != counts["rows"]:
        raise Rejected(f"{target.name} did not read back with the same rows; it was not written correctly.")
    if plain:
        pwb = load_workbook(source)
        for ws, header_row, fields in _sheets(pwb, source.name):
            fcol = next(c for c, f in fields.items() if f == "format")
            for r in _rows(ws, header_row, fields):
                for c, field in fields.items():
                    key = (ws.title, ws.cell(r, c).coordinate)
                    if key in plain_values:
                        ws.cell(r, c).value = plain_values[key]
                ws.cell(r, fcol).value = "PLAIN_TEXT"
        pwb.properties.keywords = None
        pwb.save(plain)
    return counts


def load_marks(path) -> dict:
    return read_json(path) if path else None
